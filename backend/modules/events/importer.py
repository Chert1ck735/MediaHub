"""Импорт событий из файлов.

Этот модуль отвечает только за чтение CSV/XLSX и создание объектов Event.
Он не завершает транзакции и не знает ничего про FastAPI.
"""

from __future__ import annotations

import csv
import io
import uuid
from datetime import datetime, timezone

from openpyxl import load_workbook

from backend.db.events import Event
from backend.modules.events.errors import EventImportError


class EventFileImporter:
    """Разбирает загруженный файл и превращает строки в объекты Event."""

    def parse(self, *, filename: str, content: bytes, organization_id: uuid.UUID) -> list[Event]:
        """Выбирает парсер по расширению файла и собирает список Event."""
        rows = self._rows_from_file(filename, content)
        return [self._event_from_row(row, organization_id) for row in rows]

    def _rows_from_file(self, filename: str, content: bytes) -> list[dict[str, object]]:
        """Направляет файл в CSV- или XLSX-парсер."""
        lowered = filename.lower()
        if lowered.endswith(".csv"):
            return self._rows_from_csv(content)
        if lowered.endswith(".xlsx"):
            return self._rows_from_xlsx(content)
        raise EventImportError("Поддерживаются только файлы .csv и .xlsx")

    def _rows_from_csv(self, content: bytes) -> list[dict[str, object]]:
        """Читает CSV в разных популярных кодировках и возвращает строки-словари."""
        text = None
        for encoding in ("utf-8-sig", "utf-8", "cp1251", "windows-1251", "utf-16", "latin-1"):
            try:
                text = content.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        if text is None:
            raise EventImportError(
                "Не удалось прочитать CSV. Сохраните файл в UTF-8, CP1251 или XLSX и повторите импорт."
            )
        sample = text[:2048]
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;")
        except csv.Error:
            dialect = csv.excel
            dialect.delimiter = ";"
        reader = csv.DictReader(io.StringIO(text), dialect=dialect)
        return [dict(row) for row in reader]

    def _rows_from_xlsx(self, content: bytes) -> list[dict[str, object]]:
        """Читает первый лист XLSX и превращает его строки в словари по заголовкам."""
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            return []

        headers = [str(item).strip() if item is not None else "" for item in rows[0]]
        result: list[dict[str, object]] = []
        for values in rows[1:]:
            if values is None or not any(value not in (None, "") for value in values):
                continue
            row: dict[str, object] = {}
            for index, header in enumerate(headers):
                if not header:
                    continue
                row[header] = values[index] if index < len(values) else None
            result.append(row)
        return result

    def _event_from_row(self, row: dict[str, object], organization_id: uuid.UUID) -> Event:
        """Собирает Event из одной строки импортируемого файла."""
        title = _clean(_pick(row, ("title", "name", "event", "название", "инфоповод", "тема")))
        if not title:
            raise EventImportError("У инфоповода нет названия")

        raw_date = _pick(row, ("event_date", "date", "дата", "event day", "day"))
        event_date = parse_event_date(raw_date if raw_date is not None else "")
        description = _clean(
            _pick(row, ("description", "summary", "описание", "комментарий", "details"))
        )
        location = _clean(_pick(row, ("location", "region", "место", "регион", "город")))
        external_id = _clean(_pick(row, ("external_id", "external", "url", "id")))

        return Event(
            title=title,
            description=description,
            event_date=event_date,
            location=location,
            external_id=external_id,
            imported_at=datetime.now(timezone.utc),
            organization_id=organization_id,
        )


def parse_event_date(value: str | datetime) -> datetime:
    """Приводит дату из файла или запроса к timezone-aware UTC datetime."""
    if isinstance(value, datetime):
        dt = value
    else:
        raw = str(value).strip()
        if not raw:
            raise EventImportError("Не указана дата инфоповода")

        normalized = raw.replace(".", "-").replace("/", "-")
        formats = ("%Y-%m-%d", "%d-%m-%Y", "%d-%m-%y", "%Y-%m-%d %H:%M:%S")
        dt = None
        for fmt in formats:
            try:
                dt = datetime.strptime(normalized, fmt)
                break
            except ValueError:
                continue
        if dt is None:
            try:
                dt = datetime.fromisoformat(raw)
            except ValueError as exc:
                raise EventImportError(f"Не удалось распознать дату: {raw}") from exc

    if dt.tzinfo is None:
        return dt.replace(tzinfo=timezone.utc)
    return dt.astimezone(timezone.utc)


def clean_text(value: object | None) -> str | None:
    """Публичная обертка для одинаковой очистки текстовых полей."""
    return _clean(value)


def _clean(value: object | None) -> str | None:
    """Обрезает пробелы и превращает пустую строку в None."""
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _pick(row: dict[str, object], aliases: tuple[str, ...]) -> object | None:
    """Достает значение из строки по одному из возможных названий колонки."""
    lowered = {str(key).strip().lower(): value for key, value in row.items()}
    for alias in aliases:
        if alias in lowered:
            return lowered[alias]
    return None
