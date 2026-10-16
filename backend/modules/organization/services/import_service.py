"""Импорт пользователей из Excel/CSV: preview → confirm.

Два шага намеренно: сначала показываем, что произойдёт (кого создадим, кого
обновим, где ошибки), и только по подтверждению пишем в БД.

Правила (ТЗ):
* ``email`` обязателен, нормализуется в lowercase и валидируется;
* роль по умолчанию — ``minion``;
* ``superadmin`` из файла ЗАПРЕЩЁН — иначе загрузкой файла можно поднять себе
  права до владельца всех организаций;
* группы в ``user_groups`` разделяются ``;``, отсутствующие создаются;
* ошибочная строка не отменяет остальные — импорт построчный, а не «всё или
  ничего».

Защита от DoS: файл читается в память целиком (и openpyxl, и csv), поэтому
ограничены и размер файла, и число строк. Без лимита один .xlsx на сотни
мегабайт кладёт процесс.
"""
from __future__ import annotations

import csv
import io
import re
import uuid
from typing import Iterable, Sequence

from fastapi import HTTPException, status
from openpyxl import load_workbook
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession

from backend.db.enums import UserRole, UserStatus
from backend.db.user_groups import UserGroup, UserGroupMember
from backend.db.users import User
from backend.modules.auth.services.password_service import PasswordService
from backend.modules.organization.schemas import (
    ImportPreviewResponse,
    ImportResultResponse,
    ImportRowPreview,
)
from backend.modules.organization.services.audit_service import AuditService
from backend.modules.organization.services.authorization_service import AuthorizationService

MAX_FILE_BYTES = 5 * 1024 * 1024   # 5 МБ
MAX_ROWS = 5_000

COLUMNS = ("email", "first_name", "last_name", "role", "user_groups")

# Достаточная проверка формата: строгую валидность всё равно подтвердит только
# доставка письма, а вычурная regex-«валидация» отвергает легальные адреса.
_EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")

_ROLE_ALIASES = {
    "minion": UserRole.minion,
    "admin": UserRole.admin,
    "администратор": UserRole.admin,
    "пользователь": UserRole.minion,
    "": UserRole.minion,
}


class UserImportService:
    def __init__(self, db: AsyncSession) -> None:
        self._db = db
        self._authz = AuthorizationService(db)
        self._audit = AuditService(db)
        self._passwords = PasswordService()

    # ── разбор файла ────────────────────────────────────────────────────────

    @staticmethod
    def _parse(content: bytes, filename: str) -> list[dict[str, str]]:
        """Достаёт строки из .xlsx или .csv, приводя заголовки к нижнему регистру."""
        if len(content) > MAX_FILE_BYTES:
            raise HTTPException(
                status.HTTP_413_REQUEST_ENTITY_TOO_LARGE,
                f"Файл больше {MAX_FILE_BYTES // 1024 // 1024} МБ",
            )

        name = (filename or "").lower()
        if name.endswith(".csv"):
            rows = UserImportService._parse_csv(content)
        elif name.endswith(".xlsx"):
            rows = UserImportService._parse_xlsx(content)
        else:
            raise HTTPException(
                status.HTTP_400_BAD_REQUEST,
                "Поддерживаются только файлы .xlsx и .csv",
            )

        if len(rows) > MAX_ROWS:
            raise HTTPException(
                status.HTTP_413_REQUEST_ENTITY_TOO_LARGE,
                f"В файле больше {MAX_ROWS} строк",
            )
        return rows

    @staticmethod
    def _parse_csv(content: bytes) -> list[dict[str, str]]:
        for encoding in ("utf-8-sig", "cp1251"):
            try:
                text = content.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        else:
            raise HTTPException(
                status.HTTP_400_BAD_REQUEST, "Не удалось прочитать кодировку CSV"
            )

        # Разделитель бывает и ';' (русский Excel), и ',' — определяем по первой строке.
        sample = text.split("\n", 1)[0]
        delimiter = ";" if sample.count(";") > sample.count(",") else ","

        reader = csv.DictReader(io.StringIO(text), delimiter=delimiter)
        out: list[dict[str, str]] = []
        for row in reader:
            out.append(
                {
                    (key or "").strip().lower(): (value or "").strip()
                    for key, value in row.items()
                }
            )
        return out

    @staticmethod
    def _parse_xlsx(content: bytes) -> list[dict[str, str]]:
        # read_only + data_only: не тянем стили и формулы, только значения —
        # иначе большой файл распухает в памяти в разы.
        workbook = load_workbook(
            io.BytesIO(content), read_only=True, data_only=True
        )
        sheet = workbook.active
        rows_iter = sheet.iter_rows(values_only=True)

        try:
            header = next(rows_iter)
        except StopIteration:
            return []

        headers = [str(h or "").strip().lower() for h in header]
        out: list[dict[str, str]] = []
        for values in rows_iter:
            if values is None or all(v is None or str(v).strip() == "" for v in values):
                continue  # пустая строка в хвосте листа — не ошибка
            row = {
                headers[i]: ("" if values[i] is None else str(values[i]).strip())
                for i in range(min(len(headers), len(values)))
            }
            out.append(row)
        workbook.close()
        return out

    # ── preview ─────────────────────────────────────────────────────────────

    async def preview(
        self,
        actor: User,
        organization_id: uuid.UUID,
        content: bytes,
        filename: str,
    ) -> ImportPreviewResponse:
        self._authz.require_organization_admin(actor, organization_id)

        raw_rows = self._parse(content, filename)
        if not raw_rows:
            raise HTTPException(status.HTTP_400_BAD_REQUEST, "Файл пуст")

        existing_emails = set(
            (
                await self._db.execute(
                    select(func.lower(User.email)).where(
                        User.organization_id == organization_id
                    )
                )
            ).scalars().all()
        )
        existing_groups = {
            name.strip().lower()
            for name in (
                await self._db.execute(
                    select(UserGroup.name).where(
                        UserGroup.organization_id == organization_id
                    )
                )
            ).scalars().all()
        }

        rows: list[ImportRowPreview] = []
        new_groups: set[str] = set()
        seen_emails: set[str] = set()

        for index, raw in enumerate(raw_rows, start=2):  # 1 — заголовок
            row = self._validate_row(
                index,
                raw,
                existing_emails=existing_emails,
                existing_groups=existing_groups,
                seen_emails=seen_emails,
            )
            if row.status != "error":
                for group_name in row.user_groups:
                    if group_name.strip().lower() not in existing_groups:
                        new_groups.add(group_name)
            rows.append(row)

        return ImportPreviewResponse(
            rows=rows,
            to_create=sum(1 for r in rows if r.status == "create"),
            to_update=sum(1 for r in rows if r.status == "update"),
            skipped=sum(1 for r in rows if r.status == "skip"),
            errors=sum(1 for r in rows if r.status == "error"),
            new_groups=sorted(new_groups),
        )

    def _validate_row(
        self,
        row_number: int,
        raw: dict[str, str],
        *,
        existing_emails: set[str],
        existing_groups: set[str],
        seen_emails: set[str],
    ) -> ImportRowPreview:
        errors: list[str] = []
        warnings: list[str] = []

        email = (raw.get("email") or "").strip().lower()
        if not email:
            errors.append("Не указан email")
        elif not _EMAIL_RE.match(email):
            errors.append(f"Некорректный email: {email}")
        elif email in seen_emails:
            # Дубль внутри файла: вторая строка проиграет первой — не пишем обе.
            errors.append("Дубликат email внутри файла")

        raw_role = (raw.get("role") or "").strip().lower()
        role = _ROLE_ALIASES.get(raw_role)
        if role is None:
            if raw_role in {"superadmin", "суперадмин"}:
                # Не «предупреждение», а ошибка: молча понизить роль — значит
                # сделать не то, что просил файл, и не сказать об этом.
                errors.append("Роль superadmin из файла запрещена")
            else:
                errors.append(f"Неизвестная роль: {raw_role}")

        first_name = (raw.get("first_name") or "").strip()
        last_name = (raw.get("last_name") or "").strip()
        if not first_name:
            errors.append("Не указано имя (first_name)")

        groups_raw = (raw.get("user_groups") or "").strip()
        groups = [g.strip() for g in groups_raw.split(";") if g.strip()] if groups_raw else []
        for group_name in groups:
            if group_name.strip().lower() not in existing_groups:
                warnings.append(f"Группа «{group_name}» будет создана")

        if errors:
            return ImportRowPreview(
                row_number=row_number,
                status="error",
                email=email or None,
                first_name=first_name or None,
                last_name=last_name or None,
                user_groups=groups,
                errors=errors,
                warnings=warnings,
            )

        seen_emails.add(email)
        status_value = "update" if email in existing_emails else "create"

        return ImportRowPreview(
            row_number=row_number,
            status=status_value,
            email=email,
            first_name=first_name,
            last_name=last_name,
            role=role,
            user_groups=groups,
            errors=[],
            warnings=warnings,
        )

    # ── confirm ─────────────────────────────────────────────────────────────

    async def confirm(
        self,
        actor: User,
        organization_id: uuid.UUID,
        rows: Iterable[ImportRowPreview],
    ) -> ImportResultResponse:
        """Применяет подтверждённые строки.

        Строки, помеченные ошибкой, не применяются, но и не роняют остальные.
        ``superadmin`` перепроверяем ещё раз: preview клиентский, и подсунуть в
        confirm подделанную строку с ролью superadmin — очевидная атака.
        """
        self._authz.require_organization_admin(actor, organization_id)

        created = updated = skipped = errors = 0
        created_groups: list[str] = []
        error_rows: list[ImportRowPreview] = []

        group_cache = {
            name.strip().lower(): gid
            for gid, name in (
                await self._db.execute(
                    select(UserGroup.id, UserGroup.name).where(
                        UserGroup.organization_id == organization_id
                    )
                )
            ).all()
        }

        for row in rows:
            if row.status == "error" or not row.email:
                errors += 1
                error_rows.append(row)
                continue
            if row.status == "skip":
                skipped += 1
                continue

            role = row.role or UserRole.minion
            if role == UserRole.superadmin:
                errors += 1
                error_rows.append(
                    row.model_copy(
                        update={
                            "status": "error",
                            "errors": ["Роль superadmin из файла запрещена"],
                        }
                    )
                )
                continue

            email = row.email.strip().lower()
            user = (
                await self._db.execute(
                    select(User).where(
                        func.lower(User.email) == email,
                        User.organization_id == organization_id,
                    )
                )
            ).scalar_one_or_none()

            if user is None:
                user = User(
                    email=email,
                    first_name=row.first_name or "Пользователь",
                    last_name=row.last_name or "",
                    role=role,
                    status=UserStatus.invited,
                    organization_id=organization_id,
                    password_hash=self._passwords.hash_password(uuid.uuid4().hex),
                    created_by=actor.id,
                )
                self._db.add(user)
                await self._db.flush()
                created += 1
            else:
                user.first_name = row.first_name or user.first_name
                user.last_name = row.last_name if row.last_name is not None else user.last_name
                user.role = role
                user.updated_by = actor.id
                updated += 1

            # Отсутствующие группы создаём на лету — требование ТЗ.
            for group_name in row.user_groups:
                key = group_name.strip().lower()
                group_id = group_cache.get(key)
                if group_id is None:
                    group = UserGroup(
                        organization_id=organization_id, name=group_name.strip()
                    )
                    self._db.add(group)
                    await self._db.flush()
                    group_cache[key] = group.id
                    group_id = group.id
                    created_groups.append(group.name)

                exists = await self._db.get(UserGroupMember, (group_id, user.id))
                if exists is None:
                    self._db.add(
                        UserGroupMember(
                            user_group_id=group_id,
                            user_id=user.id,
                            organization_id=organization_id,
                        )
                    )
            await self._db.flush()

        # В аудит — только сводка. Содержимое файла туда не пишем.
        self._audit.record(
            organization_id=organization_id,
            actor_user_id=actor.id,
            entity_type="user_import",
            action="import",
            after={
                "created": created,
                "updated": updated,
                "skipped": skipped,
                "errors": errors,
                "created_groups": created_groups,
            },
        )
        await self._db.commit()

        return ImportResultResponse(
            created=created,
            updated=updated,
            skipped=skipped,
            errors=errors,
            created_groups=created_groups,
            error_rows=error_rows,
        )

    # ── отчёт об ошибках ────────────────────────────────────────────────────

    @staticmethod
    def build_error_report(rows: Iterable[ImportRowPreview]) -> bytes:
        """CSV-отчёт по ошибочным строкам — то, что пользователь чинит и грузит заново."""
        buffer = io.StringIO()
        writer = csv.writer(buffer, delimiter=";")
        writer.writerow(["Строка", *COLUMNS, "Ошибки"])
        for row in rows:
            if row.status != "error":
                continue
            writer.writerow(
                [
                    row.row_number,
                    row.email or "",
                    row.first_name or "",
                    row.last_name or "",
                    row.role.value if row.role else "",
                    ";".join(row.user_groups),
                    " | ".join(row.errors),
                ]
            )
        # BOM — иначе Excel открывает кириллицу кракозябрами.
        return buffer.getvalue().encode("utf-8-sig")

    @staticmethod
    def build_template() -> bytes:
        """Шаблон для загрузки — ссылка «Скачать шаблон» в wizard'е."""
        buffer = io.StringIO()
        writer = csv.writer(buffer, delimiter=";")
        writer.writerow(COLUMNS)
        writer.writerow(
            ["ivan@example.ru", "Иван", "Тихонов", "minion", "Модераторы;Редакторы"]
        )
        return buffer.getvalue().encode("utf-8-sig")

    # ── экспорт ─────────────────────────────────────────────────────────────

    @staticmethod
    def build_export(users: Sequence[User], *, fmt: str = "csv") -> bytes:
        """Выгрузка таблицы пользователей.

        Колонки — РОВНО те же, что принимает импорт, поэтому цикл
        «выгрузил → поправил в Excel → загрузил обратно» замыкается без ручной
        перекладки столбцов. ``status`` добавлен справа как справочный: импорт
        его игнорирует, и это осознанно (статусом управляют кнопки
        блокировки, а не файл).

        Пароли и хеши не выгружаются никогда.
        """
        rows = [
            [
                user.email,
                user.first_name,
                user.last_name,
                user.role.value,
                ";".join(group.name for group in user.groups),
                user.status.value,
            ]
            for user in users
        ]
        header = [*COLUMNS, "status"]

        if fmt == "xlsx":
            from openpyxl import Workbook

            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Пользователи"
            sheet.append(header)
            for row in rows:
                sheet.append(row)
            buffer = io.BytesIO()
            workbook.save(buffer)
            return buffer.getvalue()

        buffer = io.StringIO()
        writer = csv.writer(buffer, delimiter=";")
        writer.writerow(header)
        writer.writerows(rows)
        # BOM — иначе Excel открывает кириллицу кракозябрами.
        return buffer.getvalue().encode("utf-8-sig")
