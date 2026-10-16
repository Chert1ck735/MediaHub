"""Создание организаций и экспорт пользователей."""
from __future__ import annotations

import io
import uuid

import pytest
from openpyxl import load_workbook

from backend.db import Organization, User, UserGroup
from backend.db.enums import OrgType, UserRole, UserStatus
from backend.modules.organization.services.authorization_service import AuthorizationError
from backend.modules.organization.services.group_service import UserGroupService
from backend.modules.organization.services.import_service import UserImportService
from backend.modules.organization.services.organization_service import OrganizationService
from backend.modules.organization.services.user_service import UserService

pytestmark = [
    pytest.mark.integration,
    pytest.mark.asyncio(loop_scope="session"),
]


async def _org(session, name: str = "Орг") -> Organization:
    org = Organization(name=f"{name}-{uuid.uuid4()}", type=OrgType.government)
    session.add(org)
    await session.flush()
    return org


# Заметный маркер вместо "x": короткий хеш дал бы ложное срабатывание проверки
# «хеш не утёк в выгрузку» — буква x встречается в почте вида @x.ru.
FAKE_HASH = "$2b$12$FAKEHASHSENTINEL0000000000"


async def _user(session, org: Organization, role: UserRole) -> User:
    user = User(
        first_name="Имя",
        last_name="Фамилия",
        email=f"{uuid.uuid4()}@x.ru",
        password_hash=FAKE_HASH,
        role=role,
        status=UserStatus.active,
        organization_id=org.id,
    )
    session.add(user)
    await session.flush()
    return user


# ── создание организаций ────────────────────────────────────────────────────

async def test_admin_cannot_create_organization(db_session):
    """Обычный админ не плодит организации — 403."""
    org = await _org(db_session)
    admin = await _user(db_session, org, UserRole.admin)
    await db_session.commit()

    with pytest.raises(AuthorizationError) as exc:
        await OrganizationService(db_session).create(
            admin, name="Чужая", org_type=OrgType.government, description=None
        )
    assert exc.value.status_code == 403


async def test_superadmin_creates_organization_with_admin(db_session):
    """Суперадмин заводит организацию и сразу её администратора."""
    from sqlalchemy import select

    home = await _org(db_session)
    root = await _user(db_session, home, UserRole.superadmin)
    await db_session.commit()

    created = await OrganizationService(db_session).create(
        root,
        name=f"Новая-{uuid.uuid4()}",
        org_type=OrgType.government,
        description="Тестовая",
        admin_email="Chief@X.RU",
        admin_first_name="Пётр",
        admin_last_name="Главный",
    )

    assert created.is_active is True

    admin = (
        await db_session.execute(
            select(User).where(User.organization_id == created.id)
        )
    ).scalar_one()
    assert admin.role == UserRole.admin
    # Письма не шлём — админ заводится приглашённым.
    assert admin.status == UserStatus.invited
    # email нормализован, иначе упал бы CHECK в БД.
    assert admin.email == "chief@x.ru"
    # Пароль случайный, а не пустой и не предсказуемый.
    assert admin.password_hash and admin.password_hash != FAKE_HASH


async def test_organization_without_admin_is_allowed(db_session):
    """Админа можно не создавать — организацией управляет суперадмин."""
    home = await _org(db_session)
    root = await _user(db_session, home, UserRole.superadmin)
    await db_session.commit()

    created = await OrganizationService(db_session).create(
        root, name=f"Пустая-{uuid.uuid4()}", org_type=OrgType.other, description=None
    )
    assert created.id is not None


async def test_admin_sees_only_own_organization_in_list(db_session):
    org_a = await _org(db_session, "A")
    org_b = await _org(db_session, "B")
    admin_a = await _user(db_session, org_a, UserRole.admin)
    root = await _user(db_session, org_b, UserRole.superadmin)
    await db_session.commit()

    service = OrganizationService(db_session)

    own = await service.list_organizations(admin_a)
    assert [o.id for o, _, _ in own] == [org_a.id]

    # Суперадмин видит обе.
    all_orgs = await service.list_organizations(root)
    ids = {o.id for o, _, _ in all_orgs}
    assert org_a.id in ids and org_b.id in ids


async def test_admin_cannot_deactivate_organization(db_session):
    """Отключить организацию = отрезать всех её людей. Только суперадмин."""
    org = await _org(db_session)
    admin = await _user(db_session, org, UserRole.admin)
    await db_session.commit()

    service = OrganizationService(db_session)

    # Переименовать свою — можно.
    renamed = await service.update(admin, org.id, name="Переименована")
    assert renamed.name == "Переименована"

    # Деактивировать — нет.
    with pytest.raises(AuthorizationError):
        await service.update(admin, org.id, is_active=False)


# ── экспорт ─────────────────────────────────────────────────────────────────

async def test_export_columns_match_import(db_session):
    """Экспорт даёт ровно те колонки, что принимает импорт.

    Это то, ради чего экспорт и нужен: выгрузил → поправил в Excel → загрузил
    обратно, без ручной перекладки столбцов.
    """
    org = await _org(db_session)
    admin = await _user(db_session, org, UserRole.admin)

    group = UserGroup(organization_id=org.id, name="Модераторы")
    db_session.add(group)
    await db_session.flush()
    await UserGroupService(db_session)._set_members(group, [admin.id], org.id)
    await db_session.commit()

    users = await UserService(db_session).list_for_export(org.id)
    csv_bytes = UserImportService.build_export(users, fmt="csv")
    text = csv_bytes.decode("utf-8-sig")

    header = text.splitlines()[0].split(";")
    assert header[:5] == ["email", "first_name", "last_name", "role", "user_groups"]
    assert admin.email in text
    assert "Модераторы" in text
    # Хеш пароля не должен покидать сервер ни при каких обстоятельствах.
    assert admin.password_hash not in text


async def test_export_xlsx_is_readable(db_session):
    org = await _org(db_session)
    admin = await _user(db_session, org, UserRole.admin)
    await db_session.commit()

    users = await UserService(db_session).list_for_export(org.id)
    blob = UserImportService.build_export(users, fmt="xlsx")

    sheet = load_workbook(io.BytesIO(blob)).active
    rows = list(sheet.iter_rows(values_only=True))
    assert rows[0][:5] == ("email", "first_name", "last_name", "role", "user_groups")
    assert any(row[0] == admin.email for row in rows[1:])


async def test_export_roundtrips_into_import(db_session):
    """Выгруженный файл принимается импортом без правок — цикл замкнут."""
    org = await _org(db_session)
    admin = await _user(db_session, org, UserRole.admin)
    await _user(db_session, org, UserRole.minion)
    await db_session.commit()

    users = await UserService(db_session).list_for_export(org.id)
    blob = UserImportService.build_export(users, fmt="xlsx")

    preview = await UserImportService(db_session).preview(
        admin, org.id, blob, "users.xlsx"
    )
    # Все выгруженные — уже существуют, значит это обновления, а не ошибки.
    assert preview.errors == 0
    assert preview.to_update == len(users)
    assert preview.to_create == 0
