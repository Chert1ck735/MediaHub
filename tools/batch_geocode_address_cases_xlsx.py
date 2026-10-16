from __future__ import annotations

import argparse
import json
import re
import sys
import time
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from tools.debug_comment_location import (
    _candidate_allows_locality_centroid,
    _geocode_candidate,
    _geo_matches_expected_place,
    _region_capital,
)


COORD_HEADER = "полученная координата"
LOCATION_HEADER = "название полученной локации"
ACCEPTED_STATUSES = {"hit", "area_centroid", "poi_area", "locality_centroid"}
ADMIN_PROVIDER_KINDS = {"province", "area"}


def _norm_match(value: Any) -> str:
    text = str(value or "").strip().lower().replace("ё", "е")
    for char in "\"'«».,;:()[]{}!?—–-/":
        text = text.replace(char, " ")
    return " ".join(text.split())


def _important_tokens(value: Any) -> set[str]:
    stop_words = {
        "г",
        "город",
        "край",
        "область",
        "республика",
        "район",
        "округ",
        "муниципальный",
        "муниципального",
        "област",
    }
    return {
        token
        for token in _norm_match(value).split()
        if len(token) >= 4 and token not in stop_words
    }


def _token_stem(token: str) -> str:
    token = token.strip()
    if len(token) <= 5:
        return token
    for suffix in (
        "ского",
        "ской",
        "ский",
        "ская",
        "ское",
        "ыми",
        "ого",
        "его",
        "ому",
        "ами",
        "ями",
        "ами",
        "ово",
        "ево",
        "ая",
        "ий",
        "ый",
        "ой",
        "ом",
        "ах",
        "ях",
        "ов",
        "ев",
        "ым",
        "им",
        "ы",
        "а",
        "е",
        "о",
    ):
        if token.endswith(suffix) and len(token) - len(suffix) >= 5:
            return token[: -len(suffix)]
    return token[:8]


def _tokens_loosely_subset(expected: set[str], provider: set[str]) -> bool:
    if not expected:
        return False
    provider_stems = {_token_stem(token) for token in provider}
    for token in expected:
        stem = _token_stem(token)
        if not any(
            stem == provider_stem
            or stem.startswith(provider_stem)
            or provider_stem.startswith(stem)
            for provider_stem in provider_stems
        ):
            return False
    return True


def _norm(value: Any) -> str:
    return str(value or "").strip()


def _query(region: str, case_name: str) -> str:
    if region and region.lower().replace("ё", "е") not in case_name.lower().replace("ё", "е"):
        return f"{region}, {case_name}"
    return case_name


def _looks_like_street_or_house(case_name: str) -> bool:
    low = _norm_match(case_name)
    street_markers = (
        "ул",
        "улица",
        "проспект",
        "пр кт",
        "проезд",
        "переулок",
        "шоссе",
        "тракт",
        "набережная",
        "дом",
    )
    if any(marker in low.split() for marker in street_markers):
        return True
    return bool(re.search(r"\bд(?:ом)?\s*\d", low) or re.search(r"\b[а-яё-]{5,}\s*,?\s*\d", low))


def _has_explicit_locality(case_name: str) -> bool:
    low = _norm_match(case_name)
    markers = {
        "г",
        "город",
        "село",
        "с",
        "поселок",
        "поселке",
        "пос",
        "п",
        "пгт",
        "деревня",
        "д",
        "станица",
        "аул",
    }
    return bool(markers.intersection(low.split()))


def _is_region_only_case(case_name: str) -> bool:
    low = _norm_match(case_name)
    return bool(low) and (
        low.endswith(" область")
        or low.endswith(" край")
        or low.startswith("республика ")
    ) and not any(marker in low for marker in ("район", "округ", "улица", "ул ", "дом"))


def _owner_street_variant(case_name: str) -> str:
    import re

    # "Центральный проспект Новоленда, дом 9" -> "Новоленда, Центральный проспект, дом 9".
    pattern = re.compile(
        r"^\s*(?P<street>.+?\b(?:улица|ул\.?|проспект|проезд|переулок|шоссе|тракт|набережная))"
        r"\s+(?P<owner>[А-ЯЁA-Z][А-ЯЁA-Zа-яёa-z-]{4,})"
        r"(?P<rest>\s*,?\s*(?:дом|д\.?)?\s*\d.*)$",
        re.IGNORECASE,
    )
    match = pattern.match(case_name)
    if not match:
        return ""
    owner = match.group("owner").strip(" ,.")
    street = match.group("street").strip(" ,.")
    rest = match.group("rest").strip(" ,.")
    return ", ".join(part for part in (owner, street, rest) if part)


def _case_rewrite_variants(case_name: str, *, region: str, comment_text: str) -> list[str]:
    """Turn noisy human case titles into geocoder-shaped candidates."""

    text = _norm(case_name)
    low = _norm_match(text)
    result: list[str] = []

    def add(value: str) -> None:
        value = _norm(value)
        if value and value not in result and value != text:
            result.append(value)

    normalized = re.sub(r"\bселе\b", "село", text, flags=re.IGNORECASE)
    normalized = re.sub(r"\bпос[её]лке\b", "поселок", normalized, flags=re.IGNORECASE)
    normalized = re.sub(r"\bмкр\.\s*", "микрорайон ", normalized, flags=re.IGNORECASE)
    normalized = re.sub(r"\s*-\s*", "-", normalized)
    add(normalized)

    # Frequent colloquial object forms: "мост по Перенсона" is best geocoded as the street.
    match = re.search(r"\b(?:мост|здание)\s+(?:по|на)\s+([А-ЯЁ][а-яё-]{4,})", text, re.IGNORECASE)
    if match:
        add(f"улица {match.group(1)}")

    match = re.search(r"\b(?:мост|здание)\s+([А-ЯЁ][а-яё-]{4,})", text, re.IGNORECASE)
    if match:
        add(f"улица {match.group(1)}")

    match = re.search(r"\b(?:детские сады|детский сад|садик|школа)\s+в\s+([А-ЯЁ][а-яё-]{4,})", text, re.IGNORECASE)
    if match:
        name = match.group(1).strip(" .,;")
        add(f"микрорайон {name}")
        add(name)

    match = re.search(r"\b(?:школа|детский сад|садик)\s+(?:в\s+)?(?:селе|село|с\.|с\s+)\s*([А-ЯЁ][а-яё-]{4,})", text, re.IGNORECASE)
    if match:
        add(f"село {match.group(1)}")

    match = re.search(r"\b(?:селе|село|с\.|с\s+)\s*([А-ЯЁ][а-яё-]{4,})", text, re.IGNORECASE)
    if match:
        locality = match.group(1).strip(" .,;")
        if "мкр" in low or "микрорайон" in low:
            district_match = re.search(r"\b(?:мкр\.?|микрорайон)\s*([А-ЯЁ][а-яё-]{4,})", text, re.IGNORECASE)
            if district_match:
                add(f"село {locality}, микрорайон {district_match.group(1)}")
                add(f"микрорайон {district_match.group(1)}, село {locality}")
        add(f"село {locality}")
        if locality.endswith("ов"):
            add(f"село {locality}о")
        if locality.endswith("ан"):
            add(f"село {locality}ы")

    match = re.search(r"\b(?:пгт|пос\.|пос\s+|поселок)\s*([А-ЯЁ][а-яё-]{4,})(?:\s+ул\.?\s+([А-ЯЁ][а-яё-]{4,}))?", text, re.IGNORECASE)
    if match:
        locality = match.group(1)
        street = match.group(2)
        if street:
            add(f"поселок {locality}, улица {street}")
        add(f"поселок {locality}")

    if "северо" in low and "енисейск" in low:
        add("Северо-Енисейский")
        add("городской поселок Северо-Енисейский")

    if region and _region_capital(region):
        capital = _region_capital(region)
        case_tokens = _important_tokens(text)
        comment_low = _norm_match(comment_text)
        for token in case_tokens:
            if f"микрорайон {token}" in comment_low or f"жилой район {token}" in comment_low or f"жилого района {token}" in comment_low:
                add(f"{capital}, микрорайон {text}")

    return result


def _comment_query_variants(comment_text: str, *, case_name: str, region: str) -> list[str]:
    """Extract address-like fragments from the full comment before falling back to case."""

    result: list[str] = []

    def add(value: str) -> None:
        value = _norm(value)
        if value and value not in result:
            result.append(value)

    for match in re.finditer(
        r"(?:по\s+адресу[:;]?\s*)?"
        r"(?P<place>(?:село|с\.|пгт|пос\.|поселок|деревня)\s*[А-ЯЁ][а-яё-]{3,})"
        r"[^.\n]{0,80}?"
        r"(?P<street>(?:ул\.?|улица)\s*[А-ЯЁ][а-яё-]{3,}(?:\s+[А-ЯЁ][а-яё-]{3,}){0,2})"
        r"[, ]+(?:дом|д\.?)?\s*(?P<house>\d+[А-Яа-яA-Za-z]?)",
        comment_text,
        re.IGNORECASE,
    ):
        add(f"{match.group('place')}, {match.group('street')}, {match.group('house')}")

    for match in re.finditer(
        r"(?P<street>(?:ул\.?|улица)\s*[А-ЯЁ][а-яё-]{3,}(?:\s+[А-ЯЁ][а-яё-]{3,}){0,2})"
        r"[, ]+(?:дом|д\.?)?\s*(?P<house>\d+[А-Яа-яA-Za-z]?)",
        comment_text,
        re.IGNORECASE,
    ):
        street_text = _norm_match(match.group("street"))
        if any(bad in street_text.split() for bad in {"живу", "больше", "минус", "против", "напротив", "к"}):
            continue
        add(f"{match.group('street')}, {match.group('house')}")

    if re.search(r"\bроддом\b", _norm_match(case_name)):
        match = re.search(r"\bна\s+([А-ЯЁ][а-яё-]{4,})", comment_text, re.IGNORECASE)
        if match:
            add(f"{case_name}, улица {match.group(1)}")

    if "перенсон" in _norm_match(comment_text):
        add("улица Перенсона")

    return result[:8]


def _extract_near_district_localities(comment_text: Any, case_name: str) -> list[str]:
    if not any(marker in _norm_match(case_name) for marker in ("район", "р н", "округ")):
        return []
    result: list[str] = []
    case_tokens = _important_tokens(case_name)
    for match in re.finditer(
        r"\b(?:в|во)\s+([А-ЯЁ][а-яё-]{4,})\s+([А-ЯЁа-яё-]+)\s+(?:р-н|район|округ)",
        str(comment_text or ""),
        re.IGNORECASE,
    ):
        locality = match.group(1).strip(" .,;")
        tail_tokens = _important_tokens(match.group(2))
        if tail_tokens and not tail_tokens.intersection(case_tokens):
            continue
        if locality not in result:
            result.append(locality)
    return result


def _extract_context_parts(*values: Any) -> list[str]:
    blob = "\n".join(str(value or "") for value in values if value)
    if not blob:
        return []

    patterns = [
        r"\b(?:г\.?|город(?:е|а|у|ом)?)\s*([А-ЯЁ][а-яё-]{3,})",
        r"\b([А-ЯЁ][а-яё-]+ского\s+муниципального\s+округа)",
        r"\b([А-ЯЁ][а-яё-]+ский\s+муниципальный\s+округ)",
        r"\b([А-ЯЁ][а-яё-]+ского\s+района)",
        r"\b([А-ЯЁ][а-яё-]+ский\s+район)",
        r"\b([А-ЯЁ][а-яё-]+ской\s+области)",
        r"\b([А-ЯЁ][а-яё-]+ская\s+область)",
        r"\b([А-ЯЁ][а-яё-]+ского\s+края)",
        r"\b([А-ЯЁ][а-яё-]+ский\s+край)",
        r"\b(Республика\s+[А-ЯЁ][а-яё-]+)",
    ]
    blocked = {"россии", "москве", "подмосковье", "беларуси"}
    result: list[str] = []
    seen: set[str] = set()
    for pattern in patterns:
        for match in re.finditer(pattern, blob):
            text = match.group(1).strip(" .,;:!?()[]{}«»")
            key = _norm_match(text)
            if not key or key in seen or key in blocked:
                continue
            seen.add(key)
            result.append(text)
            if len(result) >= 6:
                return result
    return result


def _infer_effective_region(region: str, *values: Any) -> str:
    if _norm(region):
        return _norm(region)
    blob = _norm_match(" ".join(str(value or "") for value in values if value))
    if "красноярск" in blob or "котюков" in blob:
        return "Красноярский край"

    explicit_regions = [
        ("ленобласти", "Ленинградская область"),
        ("ленинградская область", "Ленинградская область"),
        ("нижегородская область", "Нижегородская область"),
        ("тверская область", "Тверская область"),
        ("саратовская область", "Саратовская область"),
        ("краснодарский край", "Краснодарский край"),
        ("республика алтай", "Республика Алтай"),
        ("калужская область", "Калужская область"),
        ("орловская область", "Орловская область"),
        ("самарская область", "Самарская область"),
        ("республика адыгея", "Республика Адыгея"),
    ]
    for marker, value in explicit_regions:
        if marker in blob:
            return value
    return ""


def _build_query_variants(
    *,
    region: str,
    case_name: str,
    comment_text: str = "",
    source_title: str = "",
    public_name: str = "",
) -> list[str]:
    region = _infer_effective_region(region, public_name, source_title, comment_text, case_name)
    variants: list[str] = []

    def add(value: str) -> None:
        value = _norm(value)
        if value and value not in variants:
            variants.append(value)

    if _is_region_only_case(case_name):
        capital = _region_capital(case_name)
        if capital:
            add(f"{case_name}, {capital}")

    for query in _comment_query_variants(comment_text, case_name=case_name, region=region):
        add(_query(region, query) if region else query)

    for rewritten in _case_rewrite_variants(case_name, region=region, comment_text=comment_text):
        add(_query(region, rewritten) if region else rewritten)

    for locality in _extract_near_district_localities(comment_text, case_name):
        add(_query(region, f"{locality}, {case_name}") if region else f"{locality}, {case_name}")

    owner_variant = _owner_street_variant(case_name)
    if region and owner_variant:
        add(_query(region, owner_variant))

    for context in _extract_context_parts(comment_text, source_title):
        context_low = _norm_match(context)
        if region and any(marker in context_low for marker in ("края", "области", "республика")):
            continue
        if region and _important_tokens(context).issubset(_important_tokens(region)):
            continue
        if context_low and context_low not in _norm_match(case_name):
            add(_query(region, f"{context}, {case_name}") if region else f"{context}, {case_name}")

    capital = _region_capital(region)
    if region and capital and _looks_like_street_or_house(case_name) and not _has_explicit_locality(case_name):
        add(f"{region}, {capital}, {case_name}")

    if region and capital and not _has_explicit_locality(case_name):
        case_tokens = _important_tokens(case_name)
        comment_low = _norm_match(comment_text)
        if case_tokens and any(
            f"район {token}" in comment_low
            or f"района {token}" in comment_low
            or f"жилой район {token}" in comment_low
            or f"жилого района {token}" in comment_low
            for token in case_tokens
        ):
            add(f"{region}, {capital}, {case_name}")

    add(_query(region, case_name))
    return variants


def _candidate_allows_admin_centroid(candidate: dict[str, Any], geo: dict[str, Any]) -> bool:
    provider_kind = _norm_match(geo.get("provider_kind"))
    if provider_kind not in ADMIN_PROVIDER_KINDS:
        return False

    value = _norm_match(candidate.get("value"))
    if not any(marker in value for marker in ("область", "край", "республика", "район", "округ")):
        return False

    expected_region = candidate.get("expected_region")
    if expected_region:
        if _important_tokens(candidate.get("raw_case") or candidate.get("value")) == _important_tokens(expected_region):
            return False
        if not _geo_matches_expected_place(candidate, geo):
            return False

    provider_tokens = _important_tokens(
        " ".join(
            [
                str(geo.get("provider_name") or ""),
                str(geo.get("provider_text") or ""),
                str(geo.get("provider_description") or ""),
                str(geo.get("municipality") or ""),
                str(geo.get("region") or ""),
            ]
        )
    )
    candidate_tokens = _important_tokens(candidate.get("raw_case") or candidate.get("value"))
    return bool(candidate_tokens) and candidate_tokens.issubset(provider_tokens)


def _candidate_allows_city_district_centroid(candidate: dict[str, Any], geo: dict[str, Any]) -> bool:
    if _norm_match(geo.get("provider_kind")) != "district":
        return False
    if not candidate.get("expected_region") or not _geo_matches_expected_place(candidate, geo):
        return False
    provider_tokens = _important_tokens(
        " ".join(
            [
                str(geo.get("provider_name") or ""),
                str(geo.get("provider_text") or ""),
                str(geo.get("provider_description") or ""),
            ]
        )
    )
    candidate_tokens = _important_tokens(candidate.get("raw_case") or candidate.get("value"))
    return bool(candidate_tokens) and candidate_tokens.issubset(provider_tokens)


def _candidate_allows_approx_centroid(candidate: dict[str, Any], geo: dict[str, Any]) -> bool:
    provider_kind = _norm_match(geo.get("provider_kind"))
    if provider_kind not in {"locality", "district", "area", "province"}:
        return False

    expected_region = candidate.get("expected_region")
    if expected_region and not _geo_matches_expected_place(candidate, geo):
        return False

    provider_tokens = _important_tokens(
        " ".join(
            [
                str(geo.get("provider_name") or ""),
                str(geo.get("provider_text") or ""),
                str(geo.get("provider_description") or ""),
                str(geo.get("locality") or ""),
                str(geo.get("municipality") or ""),
                str(geo.get("region") or ""),
            ]
        )
    )
    expected_place_tokens = _important_tokens(candidate.get("expected_region")) | _important_tokens(
        candidate.get("expected_locality")
    )
    candidate_tokens = _important_tokens(candidate.get("raw_case") or candidate.get("value")) - expected_place_tokens
    if provider_kind in {"area", "province"} and not any(
        marker in _norm_match(candidate.get("raw_case") or candidate.get("value"))
        for marker in ("район", "округ", "область", "край", "республика")
    ):
        return False
    return _tokens_loosely_subset(candidate_tokens, provider_tokens)


def _looks_like_city_only_case(value: Any) -> bool:
    low = _norm_match(value)
    tokens = low.split()
    if not tokens:
        return False
    if any(marker in tokens for marker in {"ул", "улица", "дом", "проспект", "проезд", "переулок"}):
        return False
    return any(marker in tokens for marker in {"г", "город", "село", "поселок", "пгт", "деревня", "станица"})


def _is_obvious_non_address_hit(result: dict[str, Any], geo: dict[str, Any]) -> bool:
    raw = _norm_match(result.get("raw_case") or result.get("value"))
    provider_kind = _norm_match(geo.get("provider_kind"))
    provider_blob = _norm_match(
        " ".join(
            [
                str(geo.get("provider_text") or ""),
                str(geo.get("provider_name") or ""),
                str(geo.get("provider_description") or ""),
                str(geo.get("region") or ""),
            ]
        )
    )
    if not result.get("expected_region") and "россия" not in provider_blob:
        return True
    if not result.get("expected_region") and raw in {"п лично", "лично", "макс"}:
        return True
    if raw in {"промзона", "промышленная зона"}:
        return True
    expected_region = result.get("expected_region")
    if expected_region:
        capital = _norm_match(_region_capital(expected_region))
        city_only = raw in {capital, f"город {capital}", f"г {capital}"}
        if city_only:
            return True
    if not result.get("expected_region") and any(marker in raw for marker in ("школа искусств", "детская школа искусств")):
        return True
    if _looks_like_city_only_case(result.get("raw_case") or result.get("value")) and provider_kind not in {
        "locality",
        "province",
        "area",
    }:
        return True
    return False


def _split_simple_localities(case_name: str) -> list[str]:
    import re

    region_tail = re.search(
        r"\b((?:[А-ЯЁ][а-яё-]+\s+)?(?:область|край)|Республика\s+[А-ЯЁ][а-яё-]+)\s*$",
        case_name,
    )
    if region_tail:
        case_name = case_name[: region_tail.start()].strip(" ,.;")

    parts = [part.strip(" .;") for part in case_name.split(",") if part.strip(" .;")]
    if len(parts) < 2 or len(parts) > 4:
        return []
    blocked = (
        "ул.",
        "улица",
        "дом",
        "д.",
        "район",
        "р-н",
        "округ",
        "край",
        "область",
        "школ",
        "сад",
        "больниц",
    )
    result: list[str] = []
    for part in parts:
        low = part.lower().replace("ё", "е")
        if any(ch.isdigit() for ch in low):
            return []
        if any(marker in low for marker in blocked):
            return []
        if len(low.split()) > 3:
            return []
        result.append(part)
    return result


def _load_cache(path: Path) -> dict[str, dict[str, Any]]:
    if not path.exists():
        return {}
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except Exception:
        return {}
    return data if isinstance(data, dict) else {}


def _save_cache(path: Path, cache: dict[str, dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(cache, ensure_ascii=False, indent=2, default=str), encoding="utf-8")


def _result_cells(result: dict[str, Any]) -> tuple[str, str]:
    status = result.get("status")
    if status not in ACCEPTED_STATUSES:
        if status == "low_precision" and _candidate_allows_locality_centroid(result, result.get("geo") or {}):
            status = "locality_centroid"
        elif status == "low_precision" and _candidate_allows_admin_centroid(result, result.get("geo") or {}):
            status = "admin_centroid"
        elif status == "low_precision" and _candidate_allows_city_district_centroid(result, result.get("geo") or {}):
            status = "city_district_centroid"
        elif status == "low_precision" and _candidate_allows_approx_centroid(result, result.get("geo") or {}):
            status = "approx_centroid"
        else:
            return "", ""
    geo = result.get("geo")
    if not isinstance(geo, dict):
        return "", ""
    if result.get("expected_region") and not _geo_matches_expected_place(result, geo):
        return "", ""
    if _is_obvious_non_address_hit(result, geo):
        return "", ""
    lat = geo.get("lat")
    lon = geo.get("lon")
    if lat is None or lon is None:
        coord = ""
    else:
        coord = f"{lat}, {lon}"
    location_name = (
        _norm(geo.get("provider_name"))
        or _norm(geo.get("provider_text"))
        or _norm(result.get("display_title"))
        or _norm(result.get("canonical_title"))
    )
    return coord, location_name


def _geocode_cached(
    *,
    cache: dict[str, dict[str, Any]],
    cache_path: Path,
    region: str,
    case_name: str,
    timeout: int,
) -> dict[str, Any]:
    key = json.dumps([region, case_name], ensure_ascii=False)
    if key not in cache:
        query = _query(region, case_name)
        candidate = {
            "kind": "llm_address_hierarchy" if region else "production:hierarchy.case",
            "value": query,
            "raw_case": case_name,
            "expected_region": region,
            "note": "Batch geocode from exported address cases.",
        }
        cache[key] = _geocode_candidate(candidate, timeout=timeout)
        _save_cache(cache_path, cache)
    return cache[key]


def _geocode_query_cached(
    *,
    cache: dict[str, dict[str, Any]],
    cache_path: Path,
    expected_region: str,
    raw_case: str,
    query: str,
    timeout: int,
) -> dict[str, Any]:
    key = json.dumps(["query", expected_region, raw_case, query], ensure_ascii=False)
    if key not in cache:
        candidate = {
            "kind": "llm_address_hierarchy" if expected_region else "production:hierarchy.case",
            "value": query,
            "raw_case": raw_case,
            "expected_region": expected_region,
            "note": "Batch geocode query variant from exported address cases.",
        }
        cache[key] = _geocode_candidate(candidate, timeout=timeout)
        _save_cache(cache_path, cache)
    return cache[key]


def _split_result_cells(
    *,
    cache: dict[str, dict[str, Any]],
    cache_path: Path,
    region: str,
    case_name: str,
    timeout: int,
) -> tuple[str, str]:
    import re

    region_tail = re.search(
        r"\b((?:[А-ЯЁ][а-яё-]+\s+)?(?:область|край)|Республика\s+[А-ЯЁ][а-яё-]+)\s*$",
        case_name,
    )
    split_region = region
    if not split_region and region_tail:
        split_region = region_tail.group(1)

    parts = _split_simple_localities(case_name)
    if not parts:
        return "", ""
    coords: list[str] = []
    names: list[str] = []
    for part in parts:
        result = _geocode_cached(
            cache=cache,
            cache_path=cache_path,
            region=split_region,
            case_name=part,
            timeout=timeout,
        )
        coord, name = _result_cells(result)
        if not coord:
            return "", ""
        coords.append(coord)
        names.append(name or part)
    return "; ".join(coords), "; ".join(names)


def main() -> int:
    parser = argparse.ArgumentParser(description="Batch-geocode exported address cases XLSX.")
    parser.add_argument("input", type=Path)
    parser.add_argument("--output", type=Path)
    parser.add_argument("--cache", type=Path, default=Path("data/address_cases_batch_geocode_cache.json"))
    parser.add_argument("--timeout", type=int, default=15)
    parser.add_argument("--limit", type=int, default=0, help="Limit unique geocode requests; 0 means all.")
    parser.add_argument("--progress-every", type=int, default=25)
    args = parser.parse_args()

    output = args.output or args.input.with_name(f"{args.input.stem}-geocoded{args.input.suffix}")
    workbook = load_workbook(args.input)
    sheet = workbook.active
    headers = [cell.value for cell in sheet[1]]
    header_index = {value: idx + 1 for idx, value in enumerate(headers) if value}
    required = ["hierarchy_case", "public_region"]
    missing = [name for name in required if name not in header_index]
    if missing:
        raise SystemExit(f"Missing required columns: {', '.join(missing)}")

    for header in (COORD_HEADER, LOCATION_HEADER):
        if header not in header_index:
            sheet.cell(row=1, column=sheet.max_column + 1).value = header
            header_index[header] = sheet.max_column

    case_col = header_index["hierarchy_case"]
    region_col = header_index["public_region"]
    comment_col = header_index.get("comment_text")
    source_title_col = header_index.get("source_title")
    public_name_col = header_index.get("public_name")
    coord_col = header_index[COORD_HEADER]
    location_col = header_index[LOCATION_HEADER]

    pairs: dict[str, tuple[str, str]] = {}
    for row_idx in range(2, sheet.max_row + 1):
        case_name = _norm(sheet.cell(row_idx, case_col).value)
        if not case_name:
            continue
        region = _norm(sheet.cell(row_idx, region_col).value)
        key = json.dumps([region, case_name], ensure_ascii=False)
        pairs.setdefault(key, (region, case_name))

    cache = _load_cache(args.cache)
    keys = list(pairs.keys())
    if args.limit > 0:
        keys = keys[: args.limit]

    started = time.perf_counter()
    requested = 0
    for idx, key in enumerate(keys, start=1):
        if key not in cache:
            region, case_name = pairs[key]
            _geocode_cached(
                cache=cache,
                cache_path=args.cache,
                region=region,
                case_name=case_name,
                timeout=args.timeout,
            )
            requested += 1
        if idx == 1 or idx % args.progress_every == 0 or idx == len(keys):
            elapsed = time.perf_counter() - started
            print(
                f"processed_unique={idx}/{len(keys)} "
                f"new_requests={requested} "
                f"elapsed_sec={elapsed:.1f}",
                flush=True,
            )

    hits = 0
    misses = 0
    for row_idx in range(2, sheet.max_row + 1):
        case_name = _norm(sheet.cell(row_idx, case_col).value)
        region = _norm(sheet.cell(row_idx, region_col).value)
        comment_text = _norm(sheet.cell(row_idx, comment_col).value) if comment_col else ""
        source_title = _norm(sheet.cell(row_idx, source_title_col).value) if source_title_col else ""
        public_name = _norm(sheet.cell(row_idx, public_name_col).value) if public_name_col else ""
        effective_region = _infer_effective_region(region, public_name, source_title, comment_text, case_name)
        coord = ""
        location_name = ""
        for query in _build_query_variants(
            region=effective_region,
            case_name=case_name,
            comment_text=comment_text,
            source_title=source_title,
            public_name=public_name,
        ):
            result = _geocode_query_cached(
                cache=cache,
                cache_path=args.cache,
                expected_region=effective_region,
                raw_case=case_name,
                query=query,
                timeout=args.timeout,
            )
            coord, location_name = _result_cells(result)
            if coord:
                break
        if not coord:
            coord, location_name = _split_result_cells(
                cache=cache,
                cache_path=args.cache,
                region=effective_region,
                case_name=case_name,
                timeout=args.timeout,
            )
        if coord:
            hits += 1
        else:
            misses += 1
        sheet.cell(row_idx, coord_col).value = coord
        sheet.cell(row_idx, location_col).value = location_name

    workbook.save(output)
    print(f"saved={output}")
    print(f"rows_with_coordinate={hits} rows_without_coordinate={misses}")
    print(f"cache={args.cache}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
