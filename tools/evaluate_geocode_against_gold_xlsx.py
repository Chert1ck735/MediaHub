from __future__ import annotations

import argparse
import math
import re
import sys
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from tools.batch_geocode_address_cases_xlsx import (
    _build_query_variants,
    _geocode_query_cached,
    _infer_effective_region,
    _load_cache,
    _result_cells,
    _save_cache,
    _split_result_cells,
)


def _norm(value: Any) -> str:
    return str(value or "").strip()


def _has_coord(value: Any) -> bool:
    text = _norm(value)
    return bool(text) and text not in {"-", "—"} and bool(re.search(r"\d", text))


def _coords(value: Any) -> list[tuple[float, float]]:
    result: list[tuple[float, float]] = []
    for part in _norm(value).split(";"):
        match = re.search(r"(-?\d+(?:\.\d+)?)\s*,\s*(-?\d+(?:\.\d+)?)", part)
        if match:
            result.append((float(match.group(1)), float(match.group(2))))
    return result


def _distance_km(a: tuple[float, float], b: tuple[float, float]) -> float:
    lat1, lon1 = a
    lat2, lon2 = b
    radius = 6371.0
    dlat = math.radians(lat2 - lat1)
    dlon = math.radians(lon2 - lon1)
    h = (
        math.sin(dlat / 2) ** 2
        + math.cos(math.radians(lat1))
        * math.cos(math.radians(lat2))
        * math.sin(dlon / 2) ** 2
    )
    return 2 * radius * math.asin(math.sqrt(h))


def _best_distance_km(got: Any, gold: Any) -> float | None:
    got_coords = _coords(got)
    gold_coords = _coords(gold)
    if not got_coords or not gold_coords:
        return None
    return min(_distance_km(got_coord, gold_coord) for got_coord in got_coords for gold_coord in gold_coords)


def _score(got_coord: Any, gold_coord: Any, distance_km: float | None) -> str:
    got_has = _has_coord(got_coord)
    gold_has = _has_coord(gold_coord)
    if gold_has and not got_has:
        return "miss"
    if got_has and not gold_has:
        return "false_positive"
    if not got_has and not gold_has:
        return "true_empty"
    if distance_km is None:
        return "unknown_distance"
    if distance_km <= 1:
        return "match_1km"
    if distance_km <= 5:
        return "match_5km"
    if distance_km <= 20:
        return "near_20km"
    return "wrong_far"


def _header_index(sheet) -> dict[str, int]:
    return {str(cell.value): idx + 1 for idx, cell in enumerate(sheet[1]) if cell.value}


def _ensure_col(sheet, headers: dict[str, int], title: str) -> int:
    if title not in headers:
        sheet.cell(row=1, column=sheet.max_column + 1).value = title
        headers[title] = sheet.max_column
    return headers[title]


def main() -> int:
    parser = argparse.ArgumentParser(description="Evaluate tuned geocoding against gold rows in review XLSX.")
    parser.add_argument("input", type=Path)
    parser.add_argument("--output", type=Path)
    parser.add_argument("--cache", type=Path, default=Path("data/address_cases_batch_geocode_cache.json"))
    parser.add_argument("--timeout", type=int, default=15)
    parser.add_argument("--limit", type=int, default=0)
    parser.add_argument("--progress-every", type=int, default=25)
    args = parser.parse_args()

    output = args.output or args.input.with_name(f"{args.input.stem}-tuned-check{args.input.suffix}")
    workbook = load_workbook(args.input)
    sheet = workbook["dedup_review"] if "dedup_review" in workbook.sheetnames else workbook.active
    headers = _header_index(sheet)

    required = {
        "case": "Адресный кейс из БД",
        "region": "Регионы",
        "public": "Паблики",
        "source": "Источники",
        "comment": "Примеры комментариев",
        "gold_coord": "Эталонная коондината",
        "gold_name": "Эталонное название",
    }
    missing = [title for title in required.values() if title not in headers]
    if missing:
        raise SystemExit(f"Missing required columns: {', '.join(missing)}")

    tuned_coord_col = _ensure_col(sheet, headers, "Новая координата")
    tuned_name_col = _ensure_col(sheet, headers, "Новое название")
    distance_col = _ensure_col(sheet, headers, "Расстояние до эталона, км")
    score_col = _ensure_col(sheet, headers, "Оценка")

    cache = _load_cache(args.cache)
    rows = []
    for row_idx in range(2, sheet.max_row + 1):
        gold_coord = sheet.cell(row_idx, headers[required["gold_coord"]]).value
        gold_name = sheet.cell(row_idx, headers[required["gold_name"]]).value
        if _norm(gold_coord) or _norm(gold_name):
            rows.append(row_idx)
    if args.limit > 0:
        rows = rows[: args.limit]

    stats: dict[str, int] = {}
    for idx, row_idx in enumerate(rows, start=1):
        case_name = _norm(sheet.cell(row_idx, headers[required["case"]]).value)
        region = _norm(sheet.cell(row_idx, headers[required["region"]]).value)
        public_name = _norm(sheet.cell(row_idx, headers[required["public"]]).value)
        source_title = _norm(sheet.cell(row_idx, headers[required["source"]]).value)
        comment_text = _norm(sheet.cell(row_idx, headers[required["comment"]]).value)
        gold_coord = sheet.cell(row_idx, headers[required["gold_coord"]]).value
        effective_region = _infer_effective_region(region, public_name, source_title, comment_text, case_name)

        coord = ""
        location_name = ""
        for query in _build_query_variants(
            region=effective_region,
            case_name=case_name,
            comment_text=comment_text,
            source_title=source_title,
            public_name=public_name,
        ):
            result = _geocode_query_cached(
                cache=cache,
                cache_path=args.cache,
                expected_region=effective_region,
                raw_case=case_name,
                query=query,
                timeout=args.timeout,
            )
            coord, location_name = _result_cells(result)
            if coord:
                break

        if not coord:
            coord, location_name = _split_result_cells(
                cache=cache,
                cache_path=args.cache,
                region=effective_region,
                case_name=case_name,
                timeout=args.timeout,
            )

        distance = _best_distance_km(coord, gold_coord)
        score = _score(coord, gold_coord, distance)
        stats[score] = stats.get(score, 0) + 1

        sheet.cell(row_idx, tuned_coord_col).value = coord
        sheet.cell(row_idx, tuned_name_col).value = location_name
        sheet.cell(row_idx, distance_col).value = round(distance, 3) if distance is not None else None
        sheet.cell(row_idx, score_col).value = score

        if idx == 1 or idx % args.progress_every == 0 or idx == len(rows):
            print(f"processed_gold_rows={idx}/{len(rows)} stats={stats}", flush=True)

    _save_cache(args.cache, cache)
    workbook.save(output)
    print(f"saved={output}")
    print("stats=" + ", ".join(f"{key}:{stats[key]}" for key in sorted(stats)))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
