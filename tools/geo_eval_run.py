"""Прогон эталонных адресов через геокодеры и сравнение стратегий.

Берёт размеченный XLSX (см. tools/geo_eval), для каждой строки восстанавливает
контекст паблика и гоняет адрес двумя способами:

  * baseline — один запрос «регион, кейс», как было раньше: без города паблика
    и без проверки результата;
  * strategy — текущая стратегия: варианты запросов с городом паблика плюс
    отбраковка точек из чужих городов.

Затем считает попадания в эталонную координату. Так видно, что дала стратегия
и чем провайдеры отличаются друг от друга — на одних и тех же адресах.

Запуск (ключи берутся из окружения/.env):
    python -m tools.geo_eval_run эталон.xlsx --provider yandex --limit 80
    python -m tools.geo_eval_run эталон.xlsx --provider dadata_suggest
"""
from __future__ import annotations

import argparse
import time
from collections import Counter

from openpyxl import load_workbook

from backend.modules.analytics.utils.location_geocode_strategy import (
    LocationContext,
    accept_geocode_result,
    build_geocode_queries,
    infer_effective_region,
)
from tools.geo_cache import DEFAULT_PATH, GeoCache
from tools.geo_eval import distance_km, parse_point
from tools.geo_providers import PROVIDERS

COLUMNS = {
    "case": "Адресный кейс из БД",
    "address": "entities.address",
    "public": "Паблики",
    "region": "Регионы",
    "source": "Источники",
    "comment": "Примеры комментариев",
    "reference": "Эталонная коондината",
}


def load_rows(path: str, limit: int) -> list[dict]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    rows = sheet.iter_rows(values_only=True)
    header = [str(cell or "").strip() for cell in next(rows)]
    index = {key: header.index(name) for key, name in COLUMNS.items() if name in header}

    out: list[dict] = []
    for row in rows:
        if not any(row):
            continue
        reference = parse_point(row[index["reference"]])
        if reference is None:
            continue
        out.append(
            {
                "case": str(row[index["case"]] or "").strip(),
                "address": str(row[index.get("address", 0)] or "").strip(),
                "public": str(row[index.get("public", 0)] or "").strip(),
                "region": str(row[index.get("region", 0)] or "").strip(),
                "source": str(row[index.get("source", 0)] or "").strip(),
                "comment": str(row[index.get("comment", 0)] or "").strip(),
                "reference": reference,
            }
        )
        if len(out) >= limit:
            break
    workbook.close()
    return out


def geocode_baseline(fn, row: dict, region: str, timeout: int):
    """Как было: один запрос «регион, кейс», результат принимается как есть."""
    query = f"{region}, {row['case']}" if region else row["case"]
    try:
        return fn(query, timeout)
    except Exception:  # noqa: BLE001 — сбой провайдера не роняет прогон
        return None


def geocode_strategy(fn, row: dict, context: LocationContext, timeout: int):
    """Как сейчас: варианты запросов по стратегии плюс отбраковка результата.

    Возвращает (точка, отклонённые). Отклонённые нужны, чтобы понять, теряем ли
    мы верные координаты из-за слишком строгой проверки.
    """
    region, queries = build_geocode_queries(context)
    rejected = []
    for query in queries:
        try:
            point = fn(query, timeout)
        except Exception:  # noqa: BLE001
            continue
        if not point:
            continue
        if accept_geocode_result(context, query, region, point):
            return point, rejected
        rejected.append((query, point))
    return None, rejected


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("path", help="XLSX с эталонной разметкой")
    parser.add_argument("--provider", choices=sorted(PROVIDERS), default="yandex")
    parser.add_argument("--limit", type=int, default=80, help="сколько адресов взять")
    parser.add_argument("--tolerance-km", type=float, default=3.0)
    parser.add_argument("--timeout", type=int, default=15)
    parser.add_argument("--pause", type=float, default=0.15, help="пауза между запросами")
    parser.add_argument("--public-contains", help="только паблики, содержащие подстроку")
    parser.add_argument("--force-region", help="подставить регион строкам, где он пуст")
    parser.add_argument("--cache", default=str(DEFAULT_PATH), help="файл кэша ответов провайдера")
    parser.add_argument("--no-cache", action="store_true", help="ходить к провайдеру всегда")
    args = parser.parse_args()

    cache = None if args.no_cache else GeoCache(args.cache)
    fn = PROVIDERS[args.provider]
    if cache is not None:
        fn = cache.wrap(args.provider, fn)
    rows = load_rows(args.path, args.limit)
    if args.public_contains:
        needle = args.public_contains.lower()
        rows = [r for r in rows if needle in r["public"].lower()]
    if args.force_region:
        for row in rows:
            row["region"] = row["region"] or args.force_region
    print(f"адресов с эталоном: {len(rows)}, провайдер: {args.provider}\n")

    results: dict[str, Counter] = {"baseline": Counter(), "strategy": Counter()}
    lost: list = []            # отбраковали верную точку
    rejected_wrong: list = []  # отбраковали неверную — так и надо
    not_found: list = []       # геокодер не дал ничего
    requests_made = 0

    for position, row in enumerate(rows, 1):
        context = LocationContext(
            case_name=row["case"],
            public_region=row["region"],
            public_name=row["public"],
            source_title=row["source"],
            comment_text=row["comment"],
            address=row["address"],
        )
        region = infer_effective_region(context)

        strategy_point, rejected = geocode_strategy(fn, row, context, args.timeout)
        for mode, point in (
            ("baseline", geocode_baseline(fn, row, region, args.timeout)),
            ("strategy", strategy_point),
        ):
            requests_made += 1
            if point is None:
                results[mode]["не нашёл"] += 1
                if mode == "strategy":
                    good = [
                        (q, p) for q, p in rejected
                        if distance_km((p.lat, p.lon), row["reference"]) <= args.tolerance_km
                    ]
                    if good:
                        lost.append((row["case"], row["public"], good[0][0], good[0][1]))
                    elif rejected:
                        rejected_wrong.append(row["case"])
                    else:
                        not_found.append(row["case"])
                continue
            km = distance_km((point.lat, point.lon), row["reference"])
            results[mode]["попал" if km <= args.tolerance_km else "мимо"] += 1

        time.sleep(args.pause)
        if position % 20 == 0:
            print(f"  обработано {position}/{len(rows)}")

    print()
    for mode, counter in results.items():
        total = sum(counter.values())
        found = counter["попал"] + counter["мимо"]
        accuracy = f"{counter['попал'] * 100 / found:.0f}%" if found else "—"
        print(
            f"{mode:9} попал {counter['попал']:3} ({counter['попал'] * 100 // total}%), "
            f"мимо {counter['мимо']:3}, не нашёл {counter['не нашёл']:3} "
            f"| точность среди найденных {accuracy}"
        )
    print(
        f"\nразбор «не нашёл» у стратегии: отбраковано верных {len(lost)}, "
        f"отбраковано неверных {len(rejected_wrong)}, "
        f"провайдер ничего не дал {len(not_found)}"
    )

    if lost:
        print("\nпотерянные верные точки (отбраковка сработала зря):")
        for case_name, public, _query, point in lost[:15]:
            place = point.locality or point.municipality or point.provider_text
            print(f"  {case_name[:34]:34} | паблик: {public[:22]:22} | точка: {str(place)[:28]}")

    if cache is not None:
        cache.save()
        print()
        print(cache.report())
    else:
        print(f"\nзапросов к провайдеру: примерно {requests_made} (без учёта вариантов)")


if __name__ == "__main__":
    main()
