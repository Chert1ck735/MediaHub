"""Стратифицированная выборка адресных кейсов для ручной разметки геокодинга.

Читает выгрузку `address_cases.xlsx` (её отдаёт `tools/sigma_export_api.py`),
схлопывает варианты написания одного адреса и раскладывает кейсы по трём слоям:

  * `city`    — паблик города или округа, у которого определяется локалитет;
  * `capital` — паблики Красноярска и краевых служб столицы;
  * `krai`    — краевые паблики и персоны без привязки к городу.

Слой `city` берётся крупнее остальных: подстановка города паблика работает
только там, а в прежнем эталоне таких строк не было вовсе.

Координаты не запрашиваются — выборку геокодит `tools/geo_label.py`.

Запуск:
    python -m tools.geo_sample address_cases.xlsx --out geo_sample.xlsx
"""
from __future__ import annotations

import argparse
import random
from collections import defaultdict

from openpyxl import Workbook, load_workbook

from backend.modules.analytics.utils.location_geocode_strategy import (
    LocationContext,
    build_geocode_queries,
    locality_from_public,
    norm_match,
)

LAYER_SHARES = {"city": 0.47, "capital": 0.20, "krai": 0.33}

COLUMNS = (
    "id", "слой", "паблик", "регион паблика", "город паблика",
    "адресный кейс", "entities.address", "учреждение", "упоминаний",
    "комментарий", "первый запрос", "вариантов запроса",
)


def _layer(public_name: str) -> str:
    if locality_from_public(public_name):
        return "city"
    if "красноярск" in norm_match(public_name):
        return "capital"
    return "krai"


def _context(row: dict) -> LocationContext:
    return LocationContext(
        case_name=str(row.get("hierarchy_case") or "").strip(),
        public_region=str(row.get("public_region") or ""),
        public_name=str(row.get("public_name") or ""),
        source_title=str(row.get("source_title") or ""),
        comment_text=str(row.get("comment_text") or ""),
        address=str(row.get("entities_address") or ""),
    )


def _read_rows(path: str) -> list[dict]:
    sheet = load_workbook(path, read_only=True, data_only=True).active
    stream = sheet.iter_rows(values_only=True)
    header = [str(name or "") for name in next(stream)]
    return [dict(zip(header, row)) for row in stream]


def _collapse(rows: list[dict]) -> dict[str, dict]:
    """Схлопывает написания одного адреса: «с. Рыбное» и «с.Рыбное» — один кейс."""
    merged: dict[str, dict] = {}
    for row in rows:
        case_name = str(row.get("hierarchy_case") or "").strip()
        if len(case_name) < 3:
            continue
        key = f"{row.get('public_id')}|{norm_match(case_name)}"
        item = merged.get(key)
        if item is None:
            merged[key] = {"row": row, "mentions": 1}
            continue
        item["mentions"] += 1
        # Держим самое длинное написание — в нём обычно больше контекста.
        if len(case_name) > len(str(item["row"].get("hierarchy_case") or "")):
            item["row"] = row
    return merged


def sample(path: str, *, size: int, seed: int, out: str) -> None:
    merged = _collapse(_read_rows(path))

    by_layer: dict[str, list] = defaultdict(list)
    for key, item in merged.items():
        by_layer[_layer(str(item["row"].get("public_name") or ""))].append((key, item))

    rng = random.Random(seed)
    picked: list[tuple[str, dict]] = []
    for layer, share in LAYER_SHARES.items():
        pool = sorted(by_layer[layer], key=lambda pair: pair[0])
        rng.shuffle(pool)
        want = round(size * share)
        if len(pool) < want:
            print(f"слой {layer}: доступно {len(pool)}, запрошено {want}")
        picked.extend((layer, item) for _key, item in pool[:want])

    book = Workbook()
    sheet = book.active
    sheet.title = "sample"
    sheet.append(list(COLUMNS))

    taken: dict[str, int] = defaultdict(int)
    variants_total = 0
    for index, (layer, item) in enumerate(picked, start=1):
        row = item["row"]
        public_name = str(row.get("public_name") or "")
        _region, queries = build_geocode_queries(_context(row))
        taken[layer] += 1
        variants_total += len(queries)
        sheet.append([
            f"{layer}-{index:04d}",
            layer,
            public_name,
            str(row.get("public_region") or ""),
            locality_from_public(public_name),
            str(row.get("hierarchy_case") or "").strip(),
            str(row.get("entities_address") or ""),
            str(row.get("entities_institution") or ""),
            item["mentions"],
            str(row.get("comment_text") or ""),
            queries[0] if queries else "",
            len(queries),
        ])

    book.save(out)

    print(f"кейсов после схлопывания: {len(merged)}")
    for layer in LAYER_SHARES:
        print(f"  {layer:8} доступно {len(by_layer[layer]):5}  взято {taken[layer]:4}")
    print(f"выборка: {len(picked)} кейсов, до {variants_total} запросов к геокодеру")
    print(f"файл: {out}")


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("path", help="XLSX с адресными кейсами из выгрузки")
    parser.add_argument("--out", default="geo_sample.xlsx", help="куда сохранить выборку")
    parser.add_argument("--size", type=int, default=300, help="сколько кейсов взять")
    parser.add_argument("--seed", type=int, default=20260826, help="зерно для воспроизводимости")
    args = parser.parse_args()
    sample(args.path, size=args.size, seed=args.seed, out=args.out)


if __name__ == "__main__":
    main()
