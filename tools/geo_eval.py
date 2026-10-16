"""Точность геокодинга по эталонному датасету, размеченному вручную.

Читает XLSX с колонками «Адресный кейс из БД», «Полученная координата» и
«Эталонная коондината» (так в файле) и считает, насколько геокодер попадает
в размеченную человеком точку. В отличие от прокси-метрики «в городе паблика»
это настоящая точность: эталон проставлен глазами.

Считаются четыре исхода:
  * попал — координата в пределах допуска от эталонной;
  * мимо — координата есть, но далеко от эталона;
  * не нашёл — эталон есть, координаты нет (потеря покрытия);
  * нет эталона — строка не размечена, в метрики не идёт.

Запуск:
    python -m tools.geo_eval "путь/к/эталону.xlsx" [--tolerance-km 3]
"""
from __future__ import annotations

import argparse
import math
import re
from collections import Counter

from openpyxl import load_workbook

CASE_COLUMN = "Адресный кейс из БД"
GOT_COLUMN = "Полученная координата"
REF_COLUMN = "Эталонная коондината"
GOT_NAME_COLUMN = "Название от геокодера"
REF_NAME_COLUMN = "Эталонное название"

_PAIR = re.compile(r"(-?\d+[.,]\d+)\s*,\s*(-?\d+[.,]\d+)")


def parse_point(value) -> tuple[float, float] | None:
    """Первая пара координат из ячейки; в файле их бывает несколько через `;`."""
    match = _PAIR.search(str(value or ""))
    if not match:
        return None
    return float(match.group(1).replace(",", ".")), float(match.group(2).replace(",", "."))


def distance_km(first: tuple[float, float], second: tuple[float, float]) -> float:
    lat1, lon1 = math.radians(first[0]), math.radians(first[1])
    lat2, lon2 = math.radians(second[0]), math.radians(second[1])
    haversine = (
        math.sin((lat2 - lat1) / 2) ** 2
        + math.cos(lat1) * math.cos(lat2) * math.sin((lon2 - lon1) / 2) ** 2
    )
    return 6371.0 * 2 * math.asin(math.sqrt(haversine))


def evaluate(path: str, tolerance_km: float) -> None:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    rows = sheet.iter_rows(values_only=True)
    header = [str(cell or "").strip() for cell in next(rows)]
    index = {name: position for position, name in enumerate(header)}

    missing = [c for c in (CASE_COLUMN, GOT_COLUMN, REF_COLUMN) if c not in index]
    if missing:
        raise SystemExit(f"в файле нет колонок: {', '.join(missing)}")

    outcome: Counter = Counter()
    distances: list[float] = []
    worst: list[tuple[float, str, str, str]] = []

    for row in rows:
        if not any(row):
            continue
        reference = parse_point(row[index[REF_COLUMN]])
        if reference is None:
            outcome["нет эталона"] += 1
            continue

        received = parse_point(row[index[GOT_COLUMN]])
        if received is None:
            outcome["не нашёл"] += 1
            continue

        km = distance_km(received, reference)
        distances.append(km)
        if km <= tolerance_km:
            outcome["попал"] += 1
        else:
            outcome["мимо"] += 1
            worst.append((
                km,
                str(row[index[CASE_COLUMN]] or "")[:36],
                str(row[index.get(GOT_NAME_COLUMN, 0)] or "")[:30],
                str(row[index.get(REF_NAME_COLUMN, 0)] or "")[:30],
            ))

    workbook.close()

    scored = outcome["попал"] + outcome["мимо"] + outcome["не нашёл"]
    if not scored:
        raise SystemExit("в файле нет ни одной строки с эталонной координатой")

    print(f"размечено эталоном: {scored} (не размечено: {outcome['нет эталона']})")
    print(f"допуск: {tolerance_km} км\n")
    for name in ("попал", "мимо", "не нашёл"):
        count = outcome[name]
        print(f"  {name:10} {count:4} ({count * 100 / scored:.1f}%)")

    found = outcome["попал"] + outcome["мимо"]
    if found:
        print(f"\n  точность среди найденных: {outcome['попал'] * 100 / found:.1f}%")
    if distances:
        distances.sort()
        print(f"  медиана отклонения: {distances[len(distances) // 2]:.1f} км")

    if worst:
        print("\nсамые далёкие промахи:")
        for km, case_name, got_name, ref_name in sorted(worst, reverse=True)[:10]:
            print(f"  {km:8.0f} км  {case_name:36} геокодер: {got_name:30} эталон: {ref_name}")


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("path", help="XLSX с эталонной разметкой")
    parser.add_argument("--tolerance-km", type=float, default=3.0, help="допуск попадания")
    args = parser.parse_args()
    evaluate(args.path, args.tolerance_km)


if __name__ == "__main__":
    main()
